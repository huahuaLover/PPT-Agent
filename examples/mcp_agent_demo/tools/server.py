"""FastMCP tools used by the minimal agent demo."""

import asyncio
import json
import os
from pathlib import Path
from typing import Any

from fastmcp import FastMCP
from pptx import Presentation
from pptx.util import Inches, Pt
from tavily import TavilyClient

from stdio_compat import install_if_needed


mcp = FastMCP("MCP Agent Demo Tools")
MAX_FILE_CHARS = 200_000
MAX_PATH_CHARS = 512


def get_workspace() -> Path:
    """Return the configured task workspace."""
    raw_workspace = os.environ.get("WORKSPACE")
    if not raw_workspace:
        raise RuntimeError("WORKSPACE environment variable is required")
    workspace = Path(raw_workspace).resolve()
    if not workspace.is_dir():
        raise RuntimeError(f"Workspace does not exist: {workspace}")
    return workspace


def resolve_workspace_path(raw_path: str) -> Path:
    """Resolve a path and ensure it cannot escape the task workspace."""
    if not raw_path.strip():
        raise ValueError("Path cannot be empty")
    if "\n" in raw_path or len(raw_path) > MAX_PATH_CHARS:
        raise ValueError("Expected a file path, but received long-form text")
    workspace = get_workspace()
    path = Path(raw_path)
    candidate = path.resolve() if path.is_absolute() else (workspace / path).resolve()
    if not candidate.is_relative_to(workspace):
        raise ValueError(f"Path is outside workspace: {raw_path}")
    return candidate


def relative_path(path: Path) -> str:
    """Return a stable POSIX path relative to the task workspace."""
    return path.resolve().relative_to(get_workspace()).as_posix()


@mcp.tool()
async def search_web(query: str, max_results: int = 5) -> str:
    """Search the web and return concise results with source URLs."""
    query = query.strip()
    if not query:
        raise ValueError("Search query cannot be empty")

    configured_limit = int(os.environ.get("SEARCH_MAX_RESULTS", "5"))
    effective_limit = min(max(max_results, 1), configured_limit, 10)
    api_key = os.environ.get("TAVILY_API_KEY")
    if not api_key:
        raise RuntimeError("Tavily API key was not provided by the MCP client")

    client = TavilyClient(api_key=api_key)
    response = await asyncio.to_thread(
        client.search,
        query=query,
        search_depth="basic",
        max_results=effective_limit,
        include_answer=False,
        include_raw_content=False,
    )
    results = [
        {
            "title": str(item.get("title", "")),
            "url": str(item.get("url", "")),
            "content": str(item.get("content", ""))[:1000],
        }
        for item in response.get("results", [])
    ]
    return json.dumps({"query": query, "results": results}, ensure_ascii=False)


@mcp.tool()
def read_file(path: str) -> str:
    """Read a UTF-8 text file from the task workspace."""
    file_path = resolve_workspace_path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"File does not exist: {path}")
    content = file_path.read_text(encoding="utf-8")
    if len(content) > MAX_FILE_CHARS:
        raise ValueError(f"File is too large to read: {path}")
    return content


@mcp.tool()
def write_file(path: str, content: str) -> str:
    """Write UTF-8 text to a file inside the task workspace."""
    file_path = resolve_workspace_path(path)
    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_text(content, encoding="utf-8")
    return relative_path(file_path)


def require_string(value: Any, field_name: str) -> str:
    """Validate and normalize a required string field."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field_name} must be a non-empty string")
    return value.strip()


def load_slide_spec(spec_path: Path) -> tuple[str, list[dict[str, Any]]]:
    """Load and validate a slide specification."""
    with spec_path.open(encoding="utf-8") as stream:
        data = json.load(stream)
    if not isinstance(data, dict):
        raise ValueError("Slide specification must be a JSON object")

    title = require_string(data.get("title"), "title")
    slides = data.get("slides")
    if not isinstance(slides, list) or not 1 <= len(slides) <= 19:
        raise ValueError("slides must contain between 1 and 19 content slides")

    normalized: list[dict[str, Any]] = []
    for index, slide in enumerate(slides, start=1):
        if not isinstance(slide, dict):
            raise ValueError(f"slides[{index}] must be an object")
        slide_title = require_string(slide.get("title"), f"slides[{index}].title")
        bullets = slide.get("bullets")
        if not isinstance(bullets, list) or not all(
            isinstance(item, str) for item in bullets
        ):
            raise ValueError(f"slides[{index}].bullets must be a list of strings")
        if len(bullets) > 8:
            raise ValueError(f"slides[{index}] cannot contain more than 8 bullets")
        normalized.append(
            {
                "title": slide_title,
                "bullets": [
                    item.strip()[:299] + "…"
                    if len(item.strip()) > 300
                    else item.strip()
                    for item in bullets
                    if item.strip()
                ],
            }
        )
    return title, normalized


def set_text_size(shape: Any, size: Pt) -> None:
    """Apply one font size to every paragraph and run in a text shape."""
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = size
        for run in paragraph.runs:
            run.font.size = size


@mcp.tool()
def create_pptx(
    spec_path: str,
    expected_pages: int,
    output_path: str = "result.pptx",
) -> str:
    """Create a basic 16:9 PowerPoint from a JSON slide specification."""
    source = resolve_workspace_path(spec_path)
    destination = resolve_workspace_path(output_path)
    if source.suffix.lower() != ".json" or not source.is_file():
        raise ValueError(
            f"Slide specification must be an existing JSON file: {spec_path}"
        )
    if destination.suffix.lower() != ".pptx":
        raise ValueError("output_path must end with .pptx")

    title, slides = load_slide_spec(source)
    if not 2 <= expected_pages <= 20:
        raise ValueError("expected_pages must be between 2 and 20")
    if len(slides) + 1 != expected_pages:
        raise ValueError(
            f"Slide specification produces {len(slides) + 1} pages, "
            f"but expected_pages is {expected_pages}"
        )
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    set_text_size(title_slide.shapes.title, Pt(34))
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "Generated by MCP Agent Demo"
        set_text_size(title_slide.placeholders[1], Pt(18))

    for slide_data in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_data["title"]
        set_text_size(slide.shapes.title, Pt(30))

        text_frame = slide.placeholders[1].text_frame
        text_frame.clear()
        bullets = slide_data["bullets"] or [""]
        for index, bullet in enumerate(bullets):
            paragraph = (
                text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            )
            paragraph.text = bullet
            paragraph.level = 0
            paragraph.font.size = Pt(20)

    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(destination)
    if not destination.is_file() or destination.stat().st_size == 0:
        raise RuntimeError(f"Failed to create PowerPoint: {output_path}")
    return relative_path(destination)


@mcp.tool()
def finalize(outcome: str) -> str:
    """Finish a stage. Outcome must be only an existing artifact path."""
    path = resolve_workspace_path(outcome)
    if not path.is_file():
        raise FileNotFoundError(f"Outcome file does not exist: {outcome}")
    if path.stat().st_size == 0:
        raise ValueError(f"Outcome file is empty: {outcome}")
    return relative_path(path)


if __name__ == "__main__":
    install_if_needed()
    mcp.run(transport="stdio", show_banner=False)
